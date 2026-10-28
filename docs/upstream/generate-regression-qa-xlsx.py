#!/usr/bin/env python3
"""Generate docs/upstream/manual-regression-checklist-qa.xlsx from the Arena QA plan."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("manual-regression-checklist-qa.xlsx")

SECTION_NAMES = {
    "0": "0. Post-sync gates (P0)",
    "1": "1. Customized — Sim behavior Arena changed",
    "2": "2. Exclusive — Arena-only",
    "3": "3. Other Sim features — stock regression",
}

# (section_key, area, priority, item, steps, expected)
# section_key is "0" | "1" | "2" | "3" — written as SECTION_NAMES[key]
ROWS: list[tuple[str, str, str, str, str, str]] = [
    # --- Section 0: post-sync gates only (deduped) ---
    (
        "0",
        "Post-sync gates",
        "P0",
        "Migrations applied cleanly",
        "1. Open the deploy / migrate job for the target environment.\n"
        "2. Confirm the job finished green.\n"
        "3. Scan logs for duplicate migration-number errors.",
        "Migrate job is green. No duplicate migration-number errors. App starts against the migrated DB.",
    ),
    (
        "0",
        "Post-sync gates",
        "P0",
        "Output-contract dump vs pre-sync baseline",
        "1. On version-6.1-main: bun run apps/sim/scripts/diff-block-output-contracts.ts dump --out /tmp/output-contracts-before.json\n"
        "2. On the sync branch: bun run apps/sim/scripts/diff-block-output-contracts.ts diff /tmp/output-contracts-before.json\n"
        "3. Compare removals/type changes to the sync report.",
        "No unexplained removed paths or type changes beyond what the sync report already accepts (aliases or documented N/A).",
    ),
    (
        "0",
        "Post-sync gates",
        "P0",
        "Block-registry subblock ID stability",
        "1. Run: bun run apps/sim/scripts/check-block-registry.ts <previous-product-ref>\n"
        "2. Review any meta coverage failures.",
        "Subblock ID stability check passes. Any meta coverage failures are documented, not silent.",
    ),
    (
        "0",
        "Post-sync gates",
        "P0",
        "Known type-check debt note",
        "1. Note dual Next.js NextRequest identity debt in *.integration.ts.\n"
        "2. Compare the failing file list to the previous sync.",
        "Debt is documented. File list has not grown as a new Arena regression.",
    ),
    (
        "0",
        "Post-sync gates",
        "P0",
        "Usage log Arena attribution (chat + Local Copilot)",
        "1. After payer-prefix / createdAt merge, send one deployed or workspace chat turn.\n"
        "2. Send one Local Copilot turn.\n"
        "3. Open the usage ledger / Usage dashboard.",
        "Both turns appear with Arena attribution columns intact. No missing actor/payer fields from the merge.",
    ),
    (
        "0",
        "Post-sync gates",
        "P0",
        "Integration detail pages for Arena-only ids",
        "1. Open integration detail pages for arena, bing_ads, figma, image_fusion, p2_docs, spyfu.\n"
        "2. For each, note usable meta vs empty state.",
        "Each page shows usable meta or an accepted empty state for this release (no crash / blank error shell).",
    ),
    # --- Section 1: Identity and shell ---
    (
        "1",
        "Identity and shell",
        "P0",
        "Arena branding",
        "1. Load workspace home and a workflow.\n"
        "2. Check logos, favicon, API key labels, and docs title.",
        "Arena logos/favicon present. Labels say Arena key (not Sim key). Docs title is P2 Agents Docs.",
    ),
    (
        "1",
        "Identity and shell",
        "P0",
        "Sidebar Usage; Billing admins only",
        "1. As org admin: open sidebar / settings and look for Usage and Billing.\n"
        "2. As non-admin member: repeat.",
        "Usage is visible. Billing is visible only to org admins; non-admins cannot open billing.",
    ),
    (
        "1",
        "Identity and shell",
        "P0",
        "Back to Arena agents",
        "1. From the Sim/Arena agent shell, click Back to Arena agents.",
        "Browser returns to the Arena hub agents URL.",
    ),
    (
        "1",
        "Identity and shell",
        "P0",
        "Help menu stays hidden",
        "1. Inspect workspace sidebar footer / more menu.\n"
        "2. Confirm there is no standalone Help menu entry (help lives under Settings if at all).",
        "No dedicated Help menu chrome in the sidebar. Intentionally-off Arena behaviour holds.",
    ),
    (
        "1",
        "Identity and shell",
        "P0",
        "Home empty state has no CreditsChip",
        "1. Open workspace home with an empty chat surface.\n"
        "2. Look for a CreditsChip on the empty state.",
        "CreditsChip is absent. Billing remains under Usage / org-admin paths only.",
    ),
    # --- Sign-in and OAuth ---
    (
        "1",
        "Sign-in and OAuth",
        "P0",
        "Unauthenticated traffic → Arena SSO resume",
        "1. Open an agent URL while signed out.\n"
        "2. Complete Arena SSO.\n"
        "3. Confirm return into Sim.",
        "Redirect goes to Arena SSO. After login, session resumes back into the original Sim/agent path.",
    ),
    (
        "1",
        "Sign-in and OAuth",
        "P0",
        "Hub origins credentialed APIs; legal → Arena URLs",
        "1. From an allowlisted hub origin, call a credentialed API with cookies.\n"
        "2. Open legal/privacy/terms links from the product.",
        "Hub credentialed calls succeed (CORS/cookies). Legal pages redirect to Arena URLs.",
    ),
    (
        "1",
        "Sign-in and OAuth",
        "P0",
        "Session create supersedes Starter when paid is active",
        "1. Org on Starter upgrades to a paid Arena plan (or already has paid active).\n"
        "2. Trigger session create / refresh billing entitlements.",
        "Paid plan supersedes Starter. Starter is no longer the blocking entitlement.",
    ),
    (
        "1",
        "Sign-in and OAuth",
        "P0",
        "Zoom user app and Zoom admin app connect",
        "1. Connect Zoom user OAuth.\n"
        "2. Connect Zoom admin app (admin scopes).\n"
        "3. Spot-check recordings/transcripts access if credentials allow.",
        "Both apps connect successfully. Admin scopes available for recordings/transcripts.",
    ),
    (
        "1",
        "Sign-in and OAuth",
        "P0",
        "Facebook Ads and Bing Ads OAuth connect",
        "1. Start Facebook Ads OAuth and complete consent.\n"
        "2. Start Bing Ads OAuth and complete consent.",
        "Both OAuth flows complete with expected scopes. Credentials appear as connected.",
    ),
    (
        "1",
        "Sign-in and OAuth",
        "P0",
        "HubSpot shared-portal alias resolves credential",
        "1. Use a HubSpot shared-portal alias in a workflow or credential picker.\n"
        "2. Run a simple HubSpot read.",
        "The alias resolves to the correct credential; the read succeeds against the expected portal.",
    ),
    # --- Deployed chat ---
    (
        "1",
        "Deployed chat",
        "P0",
        "Landing, welcome CTAs, theme, department",
        "1. Open a deployed agent public URL (ArenaDeployedChat).\n"
        "2. Inspect landing, welcome CTAs, theme, department label.",
        "Landing, welcome CTAs, theme, and department label all render correctly.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Streamed reply with charts, images, thinking/tool segments",
        "1. Send a message that yields a streamed reply.\n"
        "2. If the workflow enables them, confirm thinking / tool-call segments.\n"
        "3. Trigger a chart and a generated image; reuse the image if UI allows.",
        "Reply streams. Thinking/tool-call segments show when enabled. ECharts charts render. Images appear and can be reused.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Builder panel Launch opens the deployed chat",
        "1. Open a workflow deployed as chat. Confirm the right panel shows Launch (lightning) beside Run.\n"
        "2. Click Launch. Confirm the deployed chat opens with fromControlBar=true.\n"
        "3. Open a workflow that is not deployed as chat or app.\n"
        "4. Click Run on the deployed workflow and run the canvas.",
        "Launch is present only for a chat or app deployment and opens that deployment. It is absent otherwise. Run still executes the workflow in the builder and is not renamed.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Threads: create, switch, rename, reload history",
        "1. Create a new thread; send a message.\n"
        "2. Switch threads; rename one; reload the page.\n"
        "3. Confirm history returns without needing webhook/schedule.",
        "Create/switch/rename work. History persists after reload. No webhook or schedule required for history.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Golden queries load and insert",
        "1. Open a deployed agent configured with golden queries.\n"
        "2. Load the list and insert one into the composer.\n"
        "3. Send it.",
        "Golden queries load and insert into the composer; send works.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Feedback submit and feedback view",
        "1. Submit prompt/reply feedback from deployed chat.\n"
        "2. Open the feedback view.",
        "Feedback submits successfully and is visible in the feedback view.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Knowledge references survive refresh",
        "1. Get a reply that includes knowledge references.\n"
        "2. Refresh the page.",
        "Knowledge references remain on the reply after refresh.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Inactive or unauthorized chat → error, no junk log",
        "1. Note execution-log / run count for the deployment.\n"
        "2. Open an inactive or unauthorized chat URL and attempt a send.\n"
        "3. Re-check the log count.",
        "Client gets an error (403 for inactive/unauthorized). No new junk execution log row is written.",
    ),
    (
        "1",
        "Deployed chat",
        "P0",
        "Voice controls stay absent",
        "1. Inspect deployed chat chrome for voice/mic controls.",
        "No voice controls are present on deployed chat.",
    ),
    # --- Workspace chat / Local Cloud ---
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Workspace home chat streams",
        "1. On workspace home, send a chat message.\n"
        "2. Watch stream, theme, charts if applicable, and usage attribution.",
        "Reply streams. Theme/charts work when produced. Usage attribution still recorded.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Allowlisted users see Local/Cloud switch",
        "1. Sign in as a user with hasAccess (not localOnly).\n"
        "2. Check workspace home and the workflow Copilot tab.",
        "Local / Cloud switch is visible on home and on the workflow Copilot tab.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Local uses Arena Copilot; Cloud uses Mothership",
        "1. With COPILOT_API_KEY configured, send on Local.\n"
        "2. Switch to Cloud and send.",
        "Local returns an Arena Copilot (local model) reply. Cloud uses Mothership.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Local model picker persists; session memory inspector",
        "1. Change the Local model picker; reload.\n"
        "2. On a longer Local thread, open the session memory inspector if shown.",
        "Picker selection persists as the user default. Session memory inspector appears on longer Local threads.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Local-only users stay on Local without the switch",
        "1. Sign in as a localOnly allowlisted user.\n"
        "2. Open home and workflow Copilot.\n"
        "3. Send a message (even if a stale Cloud preference exists).",
        "No Local/Cloud switch. Sends stay on Local.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Copilot images/files land on Arena ledger",
        "1. From Copilot, generate an image and create a file on Local.\n"
        "2. Open Usage / ledger.",
        "Image/file usage appears on the Arena ledger for the Local turns.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Floating workflow chat streams chart/image",
        "1. Open floating workflow chat.\n"
        "2. Run a prompt that produces chart and/or image output.",
        "Run streams; chart and image output render in the floating chat.",
    ),
    (
        "1",
        "Workspace chat and Local/Cloud copilot",
        "P0",
        "Deploy-modal chat preview matches deployed agent",
        "1. Open deploy modal chat preview for an agent.\n"
        "2. Compare theming/welcome behaviour to the live deployed URL.",
        "Preview matches the deployed agent experience (theme, welcome, basic send).",
    ),
    # --- Billing and workspaces (unique §1 items; package walkthrough in §2) ---
    (
        "1",
        "Billing and workspaces",
        "P0",
        "Credits debit matches ledger after one Agent turn",
        "1. Note remaining credits / known balance.\n"
        "2. Run one Agent turn.\n"
        "3. Compare Usage dashboard debit to the ledger entry (credits-per-dollar).",
        "Dashboard debit matches the ledger entry for that turn. No unexplained extra debit.",
    ),
    (
        "1",
        "Billing and workspaces",
        "P0",
        "Cost block amount shows once (no double charge)",
        "1. Run a workflow with an Agent model cost and an explicit Cost block amount.\n"
        "2. Inspect the run’s usage rows.",
        "Explicit Cost block amount appears once. No double charge vs the Agent model cost.",
    ),
    (
        "1",
        "Billing and workspaces",
        "P0",
        "Usage-log multiplier: one scaled ledger row",
        "1. Run a metered tool that uses the usage-log multiplier.\n"
        "2. Inspect ledger rows for that run.",
        "Exactly one scaled ledger row is written — not a raw row plus a scaled duplicate.",
    ),
    (
        "1",
        "Billing and workspaces",
        "P0",
        "Signup personal workspace; /workspace fallback",
        "1. Sign up a new user (or use a fresh account path).\n"
        "2. Visit /workspace with no id.",
        "Signup creates a personal workspace. /workspace fallback lands on a valid workspace.",
    ),
    (
        "1",
        "Billing and workspaces",
        "P0",
        "Daily credit refresh off unless explicitly enabled",
        "1. Confirm env does not enable daily refresh (default Arena).\n"
        "2. Observe that daily refresh does not run.\n"
        "3. If a test env enables it, confirm it then runs.",
        "Daily refresh stays off by default. It only runs when the environment explicitly enables it.",
    ),
    # --- Workflow engine ---
    (
        "1",
        "Workflow engine",
        "P1",
        "Response inside loop continues; outside ends",
        "1. Build a loop with a Response inside; run.\n"
        "2. Build a workflow with a Response outside any loop; run.",
        "Response inside a loop continues the loop. Response outside a loop ends the workflow.",
    ),
    (
        "1",
        "Workflow engine",
        "P1",
        "Empty forEach still runs the block after the loop",
        "1. Build a forEach over an empty collection with a successor block after the loop.\n"
        "2. Run the workflow.",
        "Loop body is skipped. The block after the loop still executes.",
    ),
    (
        "1",
        "Workflow engine",
        "P1",
        "Nested loops and nested subflows save and run",
        "1. Create nested loops / nested subflows (Arena allows; upstream rejects).\n"
        "2. Save and run.",
        "Nested structures save and run successfully.",
    ),
    (
        "1",
        "Workflow engine",
        "P1",
        "Agent model picker: OpenRouter free-only; Azure Anthropic + image models",
        "1. Open Agent model picker.\n"
        "2. Inspect OpenRouter list pricing.\n"
        "3. Confirm Azure Anthropic and Arena-added image models appear where expected.",
        "OpenRouter list is free-only. Azure Anthropic and Arena image models still appear.",
    ),
    (
        "1",
        "Workflow engine",
        "P1",
        "Run logs/console show charts, images, acting user",
        "1. Run a workflow that produces chart and image payloads.\n"
        "2. Open console and logs for the run.",
        "Chart payloads, image payloads, and the acting user are visible in run logs/console.",
    ),
    # --- Extended integrations ---
    (
        "1",
        "Extended integrations",
        "P1",
        "Slack: search, user-channel lookup, and post",
        "1. Run Slack search and user-channel lookup (Arena extras).\n"
        "2. Run a normal Slack post.",
        "Extra ops and stock post each return data / succeed once.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "HubSpot: campaign op + list contacts",
        "1. Run hubspot_get_campaign_metrics (or another campaign op).\n"
        "2. Run stock hubspot_list_contacts.",
        "Both operations return data successfully.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Google Slides: template tables + stock slide action",
        "1. Run create-from-template / template tables (Arena path).\n"
        "2. Run one stock slide action.",
        "Template tables path and stock slide action each succeed once.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Zoom: recordings, transcript download, stock meeting",
        "1. Fetch account recordings and download a transcript.\n"
        "2. Run one stock meeting action.",
        "Recordings/transcript path and stock meeting action succeed.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Shopify list customers (success) + list products",
        "1. Run shopify_list_customers; assert downstream can read success.\n"
        "2. Run stock shopify_list_products.\n"
        "3. Optionally list/get another resource that previously mapped outputs.",
        "list_customers succeeds and success remains available downstream. list_products succeeds. Output mapping intact.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Gmail: advanced search, read thread, send/read",
        "1. Run advanced search and read thread.\n"
        "2. Send or read a message.",
        "Advanced search/thread and send/read each succeed once.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Google Drive: upload, AI query helpers, stock file read",
        "1. Upload a file and run an AI query helper.\n"
        "2. Run a stock file read / search.",
        "Upload, AI query helpers, and stock file read/search succeed.",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Google Ads search: results consumed as array",
        "1. Run Google Ads search (GAQL).\n"
        "2. Wire a downstream block to results (e.g. forEach or Function).",
        "Downstream consumes results as an array (not a bare object / undefined).",
    ),
    (
        "1",
        "Extended integrations",
        "P1",
        "Image Generator: variations, refs, pick prior image",
        "1. Prompt for three variations from one prompt.\n"
        "2. Use a reference image.\n"
        "3. Pick a previously generated image from the conversation.",
        "Multiple variations generate. Reference images work. Prior conversation images can be picked.",
    ),
    # --- Files ---
    (
        "1",
        "Files",
        "P1",
        "Generated images/docs open from file serve route",
        "1. Generate an image and a compiled document.\n"
        "2. Open each via the file serve route.",
        "Both open successfully from the file serve route.",
    ),
    (
        "1",
        "Files",
        "P1",
        "Org logo upload and display",
        "1. Upload an org logo.\n"
        "2. Confirm it displays in the shell/branding surfaces.",
        "Logo uploads and displays correctly.",
    ),
    (
        "1",
        "Files",
        "P1",
        "Editor file upload can pick chat images",
        "1. Generate or attach an image in chat.\n"
        "2. From the editor file upload, pick an image already in the chat.",
        "Editor can select images already present in the chat.",
    ),
    # --- Section 2: Local copilot runtime ---
    (
        "2",
        "Local copilot runtime",
        "P0",
        "Tool calls, patches, audit history persist across reload",
        "1. On Local, run a turn that performs tool calls / patches.\n"
        "2. Reload the page.\n"
        "3. Inspect tool calls, patches, and audit history.",
        "Tool calls, patches, and audit history persist across reload.",
    ),
    (
        "2",
        "Local copilot runtime",
        "P0",
        "User access and user memory apply to the next turn",
        "1. Configure user access / user memory for Local.\n"
        "2. Send a follow-up turn that should use that memory/access.",
        "Next turn respects user access and user memory.",
    ),
    # --- Arena task block ---
    (
        "2",
        "Arena task block",
        "P0",
        "Create task / with fields / sub-task / with fields",
        "1. In the Arena block, run create task.\n"
        "2. Create task with fields.\n"
        "3. Create sub-task and sub-task with fields.",
        "All four create operations succeed and return task data.",
    ),
    (
        "2",
        "Arena task block",
        "P0",
        "Add comment / add comment by task number",
        "1. Add a comment on a task.\n"
        "2. Add a comment by task number.",
        "Both comment operations succeed.",
    ),
    (
        "2",
        "Arena task block",
        "P0",
        "Fetch searched tasks / search task / simple search",
        "1. Run fetch searched tasks.\n"
        "2. Run search task.\n"
        "3. Run simple search.",
        "Each search/fetch operation returns matching task data.",
    ),
    (
        "2",
        "Arena task block",
        "P0",
        "Save summary / get meetings / get token",
        "1. Run save summary.\n"
        "2. Run get meetings.\n"
        "3. Run get token.",
        "Each operation succeeds and returns expected payload.",
    ),
    (
        "2",
        "Arena task block",
        "P0",
        "Get my tasks / get my overdue tasks",
        "1. Run get my tasks.\n"
        "2. Run get my overdue tasks.",
        "Both list operations return the user’s tasks / overdue tasks.",
    ),
    (
        "2",
        "Arena task block",
        "P0",
        "Client and project selectors load Arena data",
        "1. Open the Arena block panel.\n"
        "2. Open client and project selectors.",
        "Selectors load real Arena client/project data (not empty stubs / errors).",
    ),
    # --- Arena Development ---
    (
        "2",
        "Arena Development and generated apps",
        "P1",
        "Arena Development: edit an Arena app",
        "1. Add Arena Development block.\n"
        "2. Edit an existing Arena app.",
        "Edit succeeds and persists for the Arena app.",
    ),
    (
        "2",
        "Arena Development and generated apps",
        "P1",
        "Development: generate, edit, deploy to Vercel",
        "1. Generate an app with the Development block.\n"
        "2. Edit it.\n"
        "3. Deploy to Vercel.",
        "Generate, edit, and Vercel deploy each succeed.",
    ),
    (
        "2",
        "Arena Development and generated apps",
        "P1",
        "Skill share across workspace/org",
        "1. Publish a skill into the catalog.\n"
        "2. Use/copy it in another workspace or org.",
        "Published skill is usable from the other workspace/org.",
    ),
    # --- Exclusive blocks ---
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Bing Ads, Facebook Ads, Google Ads V1",
        "1. Confirm each block appears in the block menu.\n"
        "2. Run one NL or GAQL/query op per product.",
        "Blocks appear. One operation per product returns data.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Figma file read or design action",
        "1. Confirm Figma is in the block menu.\n"
        "2. Run a file read or design action.",
        "Block visible; operation returns data.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Image Fusion",
        "1. Confirm Image Fusion is in the menu.\n"
        "2. Fuse two or more images into one output.",
        "Fused image is produced from multiple inputs.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "P2 Docs",
        "1. Confirm P2 Docs is in the menu.\n"
        "2. Run template schema, icons, or team data once.",
        "Block visible; chosen operation returns data.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Semrush and SpyFu",
        "1. Confirm both appear in the menu.\n"
        "2. Run one SEO query on each.",
        "Each returns SEO query data once.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "LinkedIn via Unipile",
        "1. Confirm Unipile/LinkedIn block is available.\n"
        "2. Read company or profile.\n"
        "3. Exercise hosted auth / reconnect if needed.",
        "Read succeeds. Hosted auth / reconnect works when required.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Chart Generator",
        "1. Confirm Chart Generator is in the menu.\n"
        "2. Produce ECharts JSON and render in chat.",
        "Chat renderer draws the chart from the JSON output.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Skyvern create or get run",
        "1. If credentials exist, create or get a Skyvern run.\n"
        "2. Otherwise mark N/A with ticket/env note.",
        "Successful create/get, or explicit N/A with ticket reference when no credential.",
    ),
    (
        "2",
        "Exclusive blocks",
        "P1",
        "Airweave successful run",
        "1. Confirm Airweave visibility for this product build.\n"
        "2. Run once, or N/A with ticket if product does not ship it.",
        "One successful run, or N/A with ticket reference.",
    ),
    # --- Deployed-chat extras ---
    (
        "2",
        "Deployed-chat extras",
        "P0",
        "Departments, assistant assets, banners, prompt feedback, golden queries",
        "1. Configure/verify departments, assistant assets, banner messages.\n"
        "2. Exercise prompt feedback and workflow golden queries on deployed chat.",
        "All Arena-only extras render and function; golden queries work with the extras pack.",
    ),
    (
        "2",
        "Deployed-chat extras",
        "P0",
        "Thread list in left nav; theme sync with Arena shell",
        "1. Open deployed chat left nav thread list.\n"
        "2. Compare theme with the Arena shell.",
        "Thread list works. Theme syncs with the Arena shell.",
    ),
    # --- Admin / identity / ops ---
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "Arena SSO login, resume, legacy cookie clear",
        "1. Log in via Arena SSO.\n"
        "2. Resume a session into Sim.\n"
        "3. On a user with an old host-only session cookie, sign out/clear and confirm legacy cookies expire.",
        "Login and resume work. Legacy session cookies are cleared (including __Secure- / Domain variants).",
    ),
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "Custom OAuth apps at org level",
        "1. Configure a custom OAuth app at org level with allowed workspaces.\n"
        "2. Connect from an allowed workspace; deny/block from a disallowed one if applicable.",
        "Allowed workspaces can use the app. Restrictions respect org configuration.",
    ),
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "Help and support issue filed and stored",
        "1. File a help/support issue via the Arena help path.\n"
        "2. Confirm it is stored (DB / admin view).",
        "Issue is accepted and stored.",
    ),
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "App banner route and workspace embed page",
        "1. Open the app banner route inside the host frame.\n"
        "2. Open the workspace embed page inside the host frame.",
        "Both render correctly inside the host frame.",
    ),
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "Default workflows import and deploy for new user",
        "1. Provision or use a new user.\n"
        "2. Confirm default workflows import and can deploy.",
        "Default workflows import and deploy for the new user.",
    ),
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "User knowledge base populate attaches to user",
        "1. Trigger user knowledge base populate.\n"
        "2. Confirm attachment to the user.",
        "KB populate attaches to the user and is usable afterward.",
    ),
    (
        "2",
        "Admin, identity, and ops",
        "P1",
        "Arena general settings (timezone and events)",
        "1. Open /workspace/arena-general-settings (no workspace id required).\n"
        "2. Save timezone and events.\n"
        "3. Confirm chat/copilot honour the settings.\n"
        "4. Spot-check workspace /settings/general still works.",
        "Page opens without a workspace id. Timezone/events save and apply to chat and copilot. Workspace general settings unchanged.",
    ),
    # --- Billing package only Arena ---
    (
        "2",
        "Billing package (Arena-only)",
        "P0",
        "Starter provisioned for a new client org",
        "1. Provision a new client org through the Arena path.\n"
        "2. Inspect subscription / usage limit.",
        "Starter is provisioned with expected usage limit / metadata.",
    ),
    (
        "2",
        "Billing package (Arena-only)",
        "P0",
        "Pro and Max checkout (team_1950 / team_6500)",
        "1. As org admin, open upgrade.\n"
        "2. Start Pro (team_1950) and/or Max (team_6500) checkout.\n"
        "3. Confirm Starter Stripe id is not used for paid checkout.",
        "Catalog shows Pro/Max flat org pricing. Checkout ignores Starter Stripe id. Paid checkout can proceed from active Starter.",
    ),
    (
        "2",
        "Billing package (Arena-only)",
        "P0",
        "Org-admin upgrade comparison overlay",
        "1. As org admin, open the upgrade comparison overlay.\n"
        "2. Review Starter / Max / Enterprise columns (Pro hidden from comparison table).",
        "Overlay shows Arena comparison with flat org pricing copy. Pro is not a comparison column.",
    ),
    (
        "2",
        "Billing package (Arena-only)",
        "P0",
        "Seat reconcile does not fight flat org pricing",
        "1. With an Arena flat org plan, adjust seats / membership.\n"
        "2. Observe seat reconcile / pricing behaviour.",
        "Seat reconcile does not multiply or fight flat org pricing.",
    ),
    (
        "2",
        "Billing package (Arena-only)",
        "P0",
        "Usage limit blocks when exhausted; allows when not",
        "1. With credits remaining, run a billable action.\n"
        "2. Exhaust credits (or use a test org at limit) and retry.",
        "Run allowed under limit. Run blocked when credits are exhausted.",
    ),
    (
        "2",
        "Billing package (Arena-only)",
        "P0",
        "Paying supersedes Starter; client-org Starter still provisions",
        "1. Confirm a new client org still gets Starter.\n"
        "2. Complete paid checkout and verify Starter is superseded.",
        "Starter provisioning still runs for client orgs. Paying supersedes Starter.",
    ),
    # --- Section 3: Builder core ---
    (
        "3",
        "Builder core",
        "P0",
        "Create, rename, duplicate, delete; undo/redo",
        "1. Create a workflow; rename; duplicate; delete.\n"
        "2. Exercise undo and redo on canvas edits.",
        "CRUD works. Undo/redo restores prior canvas state.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Drag blocks, connect edges, editor panel, variables",
        "1. Drag blocks onto the canvas and connect edges.\n"
        "2. Open the editor panel.\n"
        "3. Use a <Block.output> variable reference.",
        "Blocks connect. Panel opens. Variables resolve.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Manual run, stop, console; logs page",
        "1. Manual run a workflow; stop mid-run if possible.\n"
        "2. Read the console.\n"
        "3. Open the Logs page for the same run.",
        "Run/stop/console work. Logs page shows the same run.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Save/reload; realtime second user",
        "1. Save and reload the workflow.\n"
        "2. Have a second user open the same workflow and make an edit.",
        "State persists after reload. Second user sees realtime updates.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Deploy; API trigger and chat trigger",
        "1. Deploy the workflow.\n"
        "2. Fire an API trigger and a chat trigger.",
        "Each trigger starts a run successfully.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Schedule trigger and webhook trigger",
        "1. Configure a schedule; wait for one fire (or advance in test env).\n"
        "2. Send one signed webhook request.",
        "Schedule fires once. Webhook accepts the signed request and starts a run.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Response body; Condition and Router branches",
        "1. Configure a Response block body and run.\n"
        "2. Configure Condition and Router; take each branch once.",
        "Response returns configured body. Condition and Router take the correct branch.",
    ),
    (
        "3",
        "Builder core",
        "P0",
        "Loop and parallel small inner graph",
        "1. Run a small loop inner graph.\n"
        "2. Run a small parallel inner graph.\n"
        "3. (Nested Arena case already covered in section 1.)",
        "Loop and parallel each run the inner graph correctly.",
    ),
    # --- Workspace surfaces ---
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Home, workflow list, sidebar navigation",
        "1. Open home, workflow list, and navigate the sidebar.",
        "All surfaces load and navigate without errors.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Knowledge: create base, upload, retrieve from Agent",
        "1. Create a knowledge base; upload a document.\n"
        "2. Retrieve it from an Agent block.",
        "Upload succeeds. Agent retrieval returns the document content/citations.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Tables: create, add row, filter",
        "1. Create a table; add a row; apply a filter.",
        "Row appears and filter returns the expected subset.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Files: upload, preview PDF/image, delete",
        "1. Upload a PDF and an image.\n"
        "2. Preview each; delete one.",
        "Upload/preview/delete succeed for PDF and image.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Skills: create and attach to agent",
        "1. Create a skill.\n"
        "2. Attach it to an agent and run once.",
        "Skill attaches and is used by the agent run.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Logs, tests, dashboards, changelog pages load",
        "1. Open logs, tests, dashboards, and changelog pages.",
        "Each page loads without error.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Integrations catalog and connected credential",
        "1. Open the integrations catalog.\n"
        "2. Open a connected credential detail.",
        "Catalog and credential detail both open.",
    ),
    (
        "3",
        "Workspace surfaces",
        "P2",
        "Recently deleted: restore one workflow",
        "1. Delete a workflow.\n"
        "2. Restore it from Recently deleted.",
        "Workflow restores and opens.",
    ),
    # --- Settings ---
    (
        "3",
        "Settings",
        "P2",
        "Account, teammates, org members, secrets, credential groups",
        "1. Open each settings page and perform a trivial save/view.",
        "Each page loads; trivial actions succeed.",
    ),
    (
        "3",
        "Settings",
        "P2",
        "Arena API keys: create, use on API trigger, revoke",
        "1. Create an Arena API key.\n"
        "2. Use it on an API trigger.\n"
        "3. Revoke it and confirm rejection.",
        "Create/use succeed. After revoke, the key is rejected.",
    ),
    (
        "3",
        "Settings",
        "P2",
        "Custom tools, MCP tools, MCP servers, sandboxes",
        "1. Open each settings surface.\n"
        "2. Spot-check one create/connect action if credentials allow.",
        "Surfaces load; spot-check action succeeds or is N/A with reason.",
    ),
    (
        "3",
        "Settings",
        "P2",
        "Permission groups and audit logs (enterprise)",
        "1. If the org has enterprise, open permission groups and audit logs.\n"
        "2. Otherwise mark N/A.",
        "Pages load for enterprise orgs; N/A otherwise.",
    ),
    (
        "3",
        "Settings",
        "P2",
        "Subscription page matches Arena plan",
        "1. Open the subscription page.\n"
        "2. Compare to the Arena plan verified in section 2 billing.",
        "Subscription page matches the Arena plan (Starter/Pro/Max).",
    ),
    # --- Agent and model ---
    (
        "3",
        "Agent and model path",
        "P0",
        "Agent block calls allowed model and returns text",
        "1. Configure an Agent block with one allowed model.\n"
        "2. Run and read the text output.",
        "Agent returns text from the allowed model.",
    ),
    (
        "3",
        "Agent and model path",
        "P0",
        "Tools the agent calls show in the trace",
        "1. Run an Agent that calls at least one tool.\n"
        "2. Inspect the run trace.",
        "Tool calls appear in the trace.",
    ),
    (
        "3",
        "Agent and model path",
        "P0",
        "Memory survives a second run in the same thread",
        "1. Run an Agent with memory in a thread.\n"
        "2. Run a second turn in the same thread that depends on memory.",
        "Second run uses memory from the first.",
    ),
    # --- Stock integrations (retargeted away from §1 overlaps) ---
    (
        "3",
        "Representative stock integrations",
        "P2",
        "Google: Calendar, Sheets, Docs (Gmail/Drive covered in §1)",
        "1. If credentials exist, run one read/write each for Calendar, Sheets, Docs.\n"
        "2. Skip Gmail/Drive here if already passed in section 1 Extended; note in Notes.",
        "Each tested Google service succeeds once, or N/A with credential note. Gmail/Drive may be N/A-if-covered-in-§1.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "Microsoft: Outlook, Teams, OneDrive",
        "1. With a test account, run one read or write each.",
        "Each succeeds once, or N/A if no credential.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "Atlassian: Jira, Confluence",
        "1. Run one Jira and one Confluence happy path.",
        "Each succeeds once, or N/A if no credential.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "Dev: GitHub, Linear (or Jira if Linear unavailable)",
        "1. Run one GitHub action.\n"
        "2. Run one Linear or Jira action.",
        "Each succeeds once, or N/A if no credential.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "Comms: Discord and Microsoft Teams (Slack covered in §1)",
        "1. Run Discord and Teams stock ops.\n"
        "2. Mark Slack N/A if already covered by section 1 Extended Slack.",
        "Discord/Teams succeed once, or N/A. Slack noted as covered in §1 when applicable.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "CRM: Salesforce stock read (HubSpot covered in §1)",
        "1. Run a Salesforce stock read if credentialed.\n"
        "2. Mark HubSpot N/A if already covered in section 1.",
        "Salesforce succeeds once, or N/A. HubSpot noted as covered in §1 when applicable.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "Data: Airtable or Notion; Postgres or built-in table",
        "1. Run one Airtable or Notion op.\n"
        "2. Run one Postgres or built-in table op.",
        "Each chosen path succeeds once, or N/A if no credential.",
    ),
    (
        "3",
        "Representative stock integrations",
        "P2",
        "AI utilities: Search, Function, API block, Evaluator",
        "1. Run one happy path each for Search, Function, API block, Evaluator.",
        "All four succeed once.",
    ),
    # --- Auth and tenancy ---
    (
        "3",
        "Auth and tenancy",
        "P0",
        "Email sign-in and sign-out; invite and accept",
        "1. Sign in with email; sign out.\n"
        "2. Invite a member and accept the invite.",
        "Sign-in/out work. Invite accept grants membership.",
    ),
    (
        "3",
        "Auth and tenancy",
        "P0",
        "Member cannot open another org’s workspace",
        "1. As a member of org A, attempt to open org B’s workspace URL.",
        "Access is denied; workspace from another org is not opened.",
    ),
    (
        "3",
        "Auth and tenancy",
        "P0",
        "Archived workspace is not selectable",
        "1. Archive a workspace (or use a known archived one).\n"
        "2. Try to select it from the workspace picker.",
        "Archived workspace is not selectable.",
    ),
    (
        "3",
        "Auth and tenancy",
        "P0",
        "API key from another workspace is rejected",
        "1. Take an API key from workspace A.\n"
        "2. Call a protected API scoped to workspace B.",
        "Request is rejected (unauthorized / wrong workspace).",
    ),
]

DEDUP_LOG: list[tuple[str, str, str]] = [
    (
        "§1 Deployed chat — Inactive/unauthorized → error, no junk log",
        "§0 Sync smoke — Inactive or unauthorized deployed chat returns 403…",
        "Same 403 / no-log behaviour; keep the product UX row.",
    ),
    (
        "§1 Deployed chat — Streamed reply (thinking/tool segments folded into Steps)",
        "§0 Sync smoke — Deployed chat: golden queries / thinking / tool-call / rename…",
        "Split into §1 send-reply + golden queries + threads rows.",
    ),
    (
        "§1 Deployed chat — Golden queries; Threads rename/reload",
        "§2 Deployed-chat extras — partial golden queries / thread list overlap",
        "Extras row keeps departments/assets/banners; query/thread CRUD stays in §1.",
    ),
    (
        "§2 Exclusive — Airweave successful run",
        "§0 Sync smoke — Airweave block visible and one successful run",
        "Single Airweave happy path / N/A.",
    ),
    (
        "§1 Extended — Shopify list customers (success) + list products",
        "§0 Sync smoke — Shopify list/get workflows still map outputs including success",
        "Canonical Shopify output contract in §1 Extended.",
    ),
    (
        "§1 Extended — Google Ads search: results as array",
        "§0 Sync smoke — Google Ads search: downstream blocks consume results as an array",
        "Canonical array-consumption row in §1.",
    ),
    (
        "§1 Billing — Cost block amount shows once",
        "§2 Exclusive — Cost block (explicit cost on a run shows in usage)",
        "One Cost-block usage assertion under §1 Billing.",
    ),
    (
        "§1 Identity — Sidebar Usage; Billing admins only",
        "§1 Billing — Org admin can open billing / non-admin cannot (standalone)",
        "Merged into the Identity/shell billing-visibility row.",
    ),
    (
        "§2 Admin — Arena general settings (timezone and events)",
        "§1 Identity — /workspace/arena-general-settings opens without workspace id",
        "Folded opens-without-workspace-id into §2 Steps.",
    ),
    (
        "§2 Billing package — Pro/Max checkout; Starter provision; supersede",
        "§1 Billing — Arena plans in upgrade; Paying supersedes Starter…",
        "Package walkthrough owns plan catalog/checkout/supersede; §1 keeps debit/multiplier/refresh.",
    ),
    (
        "§1 Local/Cloud switch rows",
        "§2 Local runtime preamble “Covered in section 1 for the chat switch”",
        "Switch behaviour only in §1; §2 keeps persist + memory only.",
    ),
    (
        "§1 Extended Slack / HubSpot / Gmail / Drive / Zoom",
        "§3 stock samples for the same services",
        "§3 retargeted (Discord/Teams/Salesforce/Calendar/Sheets/Docs) or N/A-if-covered.",
    ),
    (
        "§1/§2 product rows for Chart Generator, P2 Docs, Slides, Drive",
        "§0 Sync smoke — In-process tools after directExecution removal…",
        "Product section owns the run-once checks; §0 keeps non-UI gates only.",
    ),
]


def style_header(ws, col_count: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, col_count + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def write_checklist(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Checklist"
    headers = [
        "#",
        "Section",
        "Area",
        "Priority",
        "Item",
        "Steps to test",
        "Expected behaviour",
        "Result",
        "Notes",
    ]
    ws.append(headers)
    style_header(ws, len(headers))

    for i, (section_key, area, priority, item, steps, expected) in enumerate(ROWS, start=1):
        section = SECTION_NAMES[section_key]
        ws.append([i, section, area, priority, item, steps, expected, "", ""])

    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=9):
        for cell in row:
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    ws.row_dimensions[1].height = 22
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 90

    autosize(
        ws,
        {
            1: 5,
            2: 10,
            3: 28,
            4: 10,
            5: 42,
            6: 55,
            7: 50,
            8: 12,
            9: 24,
        },
    )

    dv = DataValidation(type="list", formula1='"Pass,Fail,N/A"', allow_blank=True)
    dv.error = "Choose Pass, Fail, or N/A"
    dv.errorTitle = "Invalid result"
    ws.add_data_validation(dv)
    dv.add(f"H2:H{ws.max_row}")


def write_run_info(wb: Workbook) -> None:
    ws = wb.create_sheet("Run info")
    ws["A1"] = "Arena vs Sim manual regression checklist (QA workbook)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Mark Result Pass / Fail / N/A. Failures in sections 1–2 are Arena regressions; "
        "section 3 is stock Sim. Duplicates from the markdown checklist were collapsed — see Dedup log."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)

    ws["A4"] = "Field"
    ws["B4"] = "Value"
    style_header(ws, 2)
    for i, field in enumerate(
        ["Date", "Environment URL", "Org", "Plan", "User role", "Tester"], start=5
    ):
        ws.cell(i, 1, field)
        ws.cell(i, 2, "")

    ws["A12"] = "Intentionally off in Arena (confirm they stay hidden)"
    ws["A12"].font = Font(bold=True)
    ws["A13"] = (
        "workspace forks, custom blocks, BYOK, chat keys, Sim Mailer, the help menu, "
        "voice on deployed chat, home empty-state CreditsChip, and the Presentation block "
        "(commented out of ARENA_CUSTOM_BLOCK_REGISTRY)."
    )
    ws["A13"].alignment = Alignment(wrap_text=True)

    ws["A15"] = "Suggested pass order"
    ws["A15"].font = Font(bold=True)
    order = [
        "1. Section 0 post-sync gates (when this is a sync candidate).",
        "2. P0 identity, SSO, deployed chat, billing package, builder run, agent model.",
        "3. P0 Local / Cloud copilot and Arena task block.",
        "4. P1 exclusive blocks that this environment has credentials for.",
        "5. P1 customized Slack, HubSpot, Zoom, Slides, Shopify, Gmail, Drive, image variations.",
        "6. P2 stock surfaces and a short integration sample (skip services already covered in §1).",
    ]
    for i, line in enumerate(order, start=16):
        ws.cell(i, 1, line)

    ws["A23"] = "Counts"
    ws["A23"].font = Font(bold=True)
    p0 = sum(1 for r in ROWS if r[2] == "P0")
    p1 = sum(1 for r in ROWS if r[2] == "P1")
    p2 = sum(1 for r in ROWS if r[2] == "P2")
    ws["A24"] = f"Kept rows: {len(ROWS)} (P0={p0}, P1={p1}, P2={p2})"
    ws["A25"] = f"Dedup removals logged: {len(DEDUP_LOG)}"
    ws["A26"] = "Source markdown: docs/upstream/manual-regression-checklist.md"

    autosize(ws, {1: 100, 2: 40})
    ws.row_dimensions[2].height = 45
    ws.row_dimensions[13].height = 60


def write_dedup_log(wb: Workbook) -> None:
    ws = wb.create_sheet("Dedup log")
    headers = ["Kept location", "Removed duplicate", "Reason"]
    ws.append(headers)
    style_header(ws, 3)
    for row in DEDUP_LOG:
        ws.append(list(row))
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = wrap
            ws.row_dimensions[cell.row].height = 48
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"
    autosize(ws, {1: 55, 2: 55, 3: 50})


def main() -> None:
    wb = Workbook()
    write_checklist(wb)
    write_run_info(wb)
    write_dedup_log(wb)
    wb.save(OUT)
    print(f"Wrote {OUT} with {len(ROWS)} checklist rows and {len(DEDUP_LOG)} dedup entries")


if __name__ == "__main__":
    main()
